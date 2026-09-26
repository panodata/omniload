"""The XLSX writer: what Excel can hold, what is spelled, and what is refused.

The reader side is covered with the other spreadsheet formats; this file is about
``write_xlsx``, read back with openpyxl as an independent decoder and through the
package's own reader where the two disagree.
"""

import csv
import datetime
import decimal
import logging
import sys

import openpyxl
import pytest

from dlt_filesystem.source.error import MissingDecoderError
from dlt_filesystem.source.fsspec.local import LocalFilesystemSource
from dlt_filesystem.target import writer
from dlt_filesystem.target.registry import registration_for_format, writer_for_format
from dlt_filesystem.target.writer import write_xlsx


def _sheet(path):
    return openpyxl.load_workbook(path).worksheets[0]


def _rows(path):
    """Header row to dicts, keeping empty cells as None."""
    header, *body = _sheet(path).iter_rows(values_only=True)
    return [dict(zip(header, row)) for row in body]


def _read_via_source(path, selector="#sheet_name=rows"):
    """Read a workbook back through the package's own reader."""
    return list(LocalFilesystemSource().dlt_source(f"file://{path}{selector}", ""))


def test_xlsx_is_registered_to_receive_the_table_name():
    registration = registration_for_format("xlsx")
    assert registration.writer is write_xlsx
    assert registration.takes_table_name
    assert not registration_for_format("csv").takes_table_name


def test_write_keeps_native_cells_and_the_column_union(tmp_path):
    path = tmp_path / "out.xlsx"
    rows = [
        {"id": 1, "name": "Zoë", "ok": True},
        {"id": 2, "name": "Ōtautahi", "ok": False, "note": "late column"},
    ]
    write_xlsx(str(path), rows, table_name="rows")

    assert _sheet(path).title == "rows"
    assert _rows(path) == [
        {"id": 1, "name": "Zoë", "ok": True, "note": None},
        {"id": 2, "name": "Ōtautahi", "ok": False, "note": "late column"},
    ]
    # `bool` subclasses `int`, so the wrong dispatch order writes `True` as `1`.
    assert [cell.data_type for cell in _sheet(path)[2]][:3] == ["n", "s", "b"]


def test_write_reads_back_through_the_package_reader(tmp_path):
    path = tmp_path / "out.xlsx"
    write_xlsx(str(path), [{"id": 1, "name": "Zoë"}], table_name="rows")

    assert _read_via_source(path) == [{"id": 1, "name": "Zoë"}]


def test_write_without_a_table_name_uses_excels_default_sheet(tmp_path):
    path = tmp_path / "out.xlsx"
    writer_for_format("xlsx")(str(path), [{"id": 1}])

    assert _sheet(path).title == "Sheet1"


@pytest.mark.parametrize(
    ("table_name", "sheet"),
    [
        ("a" * 40, "a" * 31),
        ("q1:[draft]*?/\\", "q1__draft_____"),
        ("'quoted'", "quoted"),
        ("'''", "Sheet1"),
    ],
)
def test_a_name_excel_refuses_is_shortened_with_a_warning(
    tmp_path, caplog, table_name, sheet
):
    """The sheet name is metadata, so a name past Excel's rules is changed rather than
    refused: with no `--dest-table` it is the path stem, and a long file name would
    otherwise fail the whole export. The warning names both."""
    path = tmp_path / "out.xlsx"
    with caplog.at_level(logging.WARNING, logger=writer.__name__):
        write_xlsx(str(path), [{"id": 1}], table_name=table_name)

    assert _sheet(path).title == sheet
    assert table_name in caplog.text and f"'{sheet}'" in caplog.text


def test_columns_differing_only_by_case_stay_two_columns(tmp_path):
    """No Excel table object is written. Its headers are case-insensitive, and
    `DataFrame.write_excel` always writes one, which leaves this sheet reading back
    empty: a later move to that API fails here."""
    path = tmp_path / "out.xlsx"
    write_xlsx(str(path), [{"A": 1, "a": 2}], table_name="rows")

    assert _rows(path) == [{"A": 1, "a": 2}]
    assert _read_via_source(path) == [{"A": 1, "a": 2}]
    assert not _sheet(path).tables


@pytest.mark.parametrize(
    "value", ["https://example.org/a", "=1+1", "0012", "1e5", "2020-01-01"]
)
def test_a_string_stays_a_plain_string(tmp_path, value):
    """No hyperlink, formula or number is made of a string that looks like one."""
    path = tmp_path / "out.xlsx"
    write_xlsx(str(path), [{"s": value}], table_name="rows")

    cell = _sheet(path)["A2"]
    assert (cell.value, cell.data_type, cell.hyperlink) == (value, "s", None)


def test_an_empty_string_stays_distinct_from_null(tmp_path):
    path = tmp_path / "out.xlsx"
    rows = [{"id": 1, "s": ""}, {"id": 2, "s": "x"}]
    write_xlsx(str(path), rows, table_name="rows")

    assert _rows(path) == rows
    assert _read_via_source(path) == rows


def test_a_float_carries_no_display_format(tmp_path):
    """Polars' writer gives a float a three-decimal display format."""
    path = tmp_path / "out.xlsx"
    write_xlsx(str(path), [{"f": 1.5}], table_name="rows")

    assert _sheet(path)["A2"].number_format == "General"


def test_a_float_is_written_to_sixteen_significant_digits(tmp_path):
    """xlsxwriter serializes a float with `%.16G`, so one needing all 17 digits reads
    back one step off. Excel itself holds 15, so this is documented rather than
    refused; pinned so a serializer change is noticed."""
    path = tmp_path / "out.xlsx"
    write_xlsx(str(path), [{"f": 0.1 + 0.2}, {"f": 0.5}], table_name="rows")

    assert [row["f"] for row in _rows(path)] == [0.3, 0.5]


def test_control_characters_round_trip_through_the_package_reader(tmp_path):
    """xlsxwriter escapes XML-illegal characters as `_xHHHH_` and escapes a literal
    `_x` sequence in turn, and calamine decodes both back. openpyxl leaves the escapes
    in, which is why this reads through the package's reader."""
    path = tmp_path / "out.xlsx"
    rows = [{"s": "a\x00b\x07c", "t": "_x0041_"}]
    write_xlsx(str(path), rows, table_name="rows")

    assert _read_via_source(path) == rows


# --- spelled as write_csv spells them ---


@pytest.mark.parametrize(
    "value",
    [
        [1, "p", True],
        {"k": 1, "nested": {"n": None}},
        (1, 2),
        b"hi",
        decimal.Decimal("1.50"),
    ],
    ids=["list", "dict", "tuple", "bytes", "decimal"],
)
def test_a_value_excel_has_no_cell_for_is_spelled_as_csv_spells_it(tmp_path, value):
    """One value reads the same in every flat format."""
    csv_path, xlsx_path = tmp_path / "out.csv", tmp_path / "out.xlsx"
    writer_for_format("csv")(str(csv_path), [{"v": value}])
    write_xlsx(str(xlsx_path), [{"v": value}], table_name="rows")

    with open(csv_path, newline="", encoding="utf-8") as handle:
        (as_csv,) = csv.DictReader(handle)
    assert _rows(xlsx_path) == [{"v": as_csv["v"]}]


# --- dates and times ---


def test_dates_and_times_are_date_cells(tmp_path):
    path = tmp_path / "out.xlsx"
    rows = [
        {
            "at": datetime.datetime(2020, 1, 2, 3, 4, 5),
            "at_ms": datetime.datetime(2020, 1, 2, 3, 4, 5, 123000),
            "day": datetime.date(2020, 1, 2),
            "clock": datetime.time(9, 30),
        }
    ]
    write_xlsx(str(path), rows, table_name="rows")

    cells = _sheet(path)[2]
    assert [cell.number_format for cell in cells] == [
        "yyyy-mm-dd hh:mm:ss",
        "yyyy-mm-dd hh:mm:ss.000",
        "yyyy-mm-dd",
        "hh:mm:ss",
    ]
    # openpyxl has no date-only cell type, so a date reads as its midnight.
    assert [cell.value for cell in cells] == [
        datetime.datetime(2020, 1, 2, 3, 4, 5),
        datetime.datetime(2020, 1, 2, 3, 4, 5, 123000),
        datetime.datetime(2020, 1, 2),
        datetime.time(9, 30),
    ]


def test_microseconds_read_back_at_millisecond_resolution(tmp_path):
    """The serial holds the microseconds; both readers round to the millisecond."""
    path = tmp_path / "out.xlsx"
    write_xlsx(
        str(path),
        [{"at": datetime.datetime(2020, 1, 2, 3, 4, 5, 123456)}],
        table_name="rows",
    )

    expected = datetime.datetime(2020, 1, 2, 3, 4, 5, 123000)
    assert _rows(path) == [{"at": expected}]
    assert _read_via_source(path) == [{"at": expected}]


@pytest.mark.parametrize(
    "zone",
    [datetime.timezone.utc, datetime.timezone(datetime.timedelta(hours=12))],
    ids=["utc", "fixed-offset"],
)
def test_an_aware_datetime_is_written_as_the_same_instant_in_utc(tmp_path, zone):
    path = tmp_path / "out.xlsx"
    at = datetime.datetime(2020, 1, 2, 15, 4, 5, tzinfo=zone)
    write_xlsx(str(path), [{"at": at}], table_name="rows")

    (row,) = _rows(path)
    assert row["at"] == at.astimezone(datetime.timezone.utc).replace(tzinfo=None)


def test_the_first_date_past_the_leap_year_region_is_written(tmp_path):
    path = tmp_path / "out.xlsx"
    first = datetime.datetime(1900, 3, 1)
    last = datetime.datetime(9999, 12, 31, 23, 59, 59, 999000)
    write_xlsx(str(path), [{"at": first}, {"at": last}], table_name="rows")

    assert [row["at"] for row in _rows(path)] == [first, last]
    assert [row["at"] for row in _read_via_source(path)] == [first, last]


# --- empty loads ---


@pytest.mark.parametrize("rows", [[], [{}, {}]], ids=["no-rows", "no-columns"])
def test_a_load_with_no_columns_is_one_empty_worksheet(tmp_path, rows):
    """A sheet with no cells is a valid workbook, so unlike Feather and ORC a load of
    `{}` rows is written rather than refused.

    The reader skips a blank sheet when it reads the whole workbook, and refuses one
    it is pointed at unless told `raise_if_empty=false`: that is the reader's rule for
    any blank worksheet, pinned here because it is what reading an empty export takes.
    """
    path = tmp_path / "out.xlsx"
    write_xlsx(str(path), rows, table_name="rows")

    assert list(_sheet(path).iter_rows(values_only=True)) in ([], [(None,)])
    assert _read_via_source(path, "") == []
    assert _read_via_source(path, "#sheet_name=rows&raise_if_empty=false") == []
    with pytest.raises(Exception, match="empty Excel sheet"):
        _read_via_source(path)


def test_an_interior_all_null_record_is_an_empty_row(tmp_path):
    """Written as an empty row, which the reader drops by default. Documented rather
    than refused: the row carries no value to lose, only its position."""
    path = tmp_path / "out.xlsx"
    write_xlsx(str(path), [{"id": 1}, {}, {"id": 3}], table_name="rows")

    assert _rows(path) == [{"id": 1}, {"id": None}, {"id": 3}]
    assert _read_via_source(path) == [{"id": 1}, {"id": 3}]


# --- refused ---


def _assert_refused(tmp_path, rows, match):
    path = tmp_path / "out.xlsx"
    with pytest.raises(ValueError, match=match):
        write_xlsx(str(path), rows, table_name="rows")
    assert not path.exists()
    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize("value", [2**53 + 1, -(2**53) - 1, 2**64])
def test_an_integer_a_double_cannot_hold_is_refused(tmp_path, value):
    _assert_refused(tmp_path, [{"id": 1, "n": value}], r"Column 'n'.*row 1")


def test_an_integer_a_double_holds_is_written(tmp_path):
    path = tmp_path / "out.xlsx"
    write_xlsx(str(path), [{"n": 2**53}, {"n": -(2**53)}], table_name="rows")

    assert [row["n"] for row in _rows(path)] == [2**53, -(2**53)]


@pytest.mark.parametrize("value", [float("nan"), float("inf"), float("-inf")])
def test_a_float_excel_has_no_number_for_is_refused(tmp_path, value):
    _assert_refused(tmp_path, [{"id": 1}, {"f": value}], r"Column 'f'.*row 2")


def test_a_string_past_the_cell_limit_is_refused(tmp_path):
    _assert_refused(tmp_path, [{"s": "x" * 32_768}], r"Column 's'.*32768 characters")


def test_the_longest_string_a_cell_holds_is_written(tmp_path):
    path = tmp_path / "out.xlsx"
    write_xlsx(str(path), [{"s": "x" * 32_767}], table_name="rows")

    assert _rows(path) == [{"s": "x" * 32_767}]


def test_a_value_that_spells_past_the_cell_limit_is_refused(tmp_path):
    """Checked on the cell as written: the string fits, its JSON text does not."""
    _assert_refused(tmp_path, [{"tags": ["x" * 32_767]}], r"Column 'tags'")


def test_a_column_name_past_the_cell_limit_is_refused(tmp_path):
    """The header is a string cell too."""
    name = "c" * 32_768
    _assert_refused(tmp_path, [{name: 1}], "as its name")


@pytest.mark.parametrize(
    "value",
    [
        datetime.datetime(1900, 1, 1),
        datetime.datetime(1900, 2, 28, 23, 59, 59),
        datetime.date(1900, 2, 28),
        datetime.date(1, 1, 1),
    ],
)
def test_a_date_in_excels_leap_year_region_is_refused(tmp_path, value):
    """`datetime(1900, 1, 1)` is written as serial 0 and reads back as 31 December
    1899, and every date before March 1900 sits in the region where Excel counts a
    29 February that never was."""
    _assert_refused(tmp_path, [{"at": value}], r"Column 'at'.*1900-03-01")


def test_a_datetime_that_rounds_into_year_10000_is_refused(tmp_path):
    """`datetime.max` reads back as year 10000, which the package reader panics on."""
    _assert_refused(tmp_path, [{"at": datetime.datetime.max}], r"Column 'at'")


@pytest.mark.parametrize(
    "clock",
    [datetime.time(23, 59, 59, 999500), datetime.time.max],
    ids=["half-millisecond", "max"],
)
def test_a_time_that_rounds_to_midnight_is_refused(tmp_path, clock):
    """A time cell holds the fraction of a day, so a time the readers round up to the
    next millisecond past 23:59:59.999 reads back as midnight."""
    _assert_refused(tmp_path, [{"clock": clock}], r"Column 'clock'.*midnight")


def test_a_time_just_short_of_the_rounding_boundary_is_written(tmp_path):
    """Up to the half millisecond, the readers round down to 23:59:59.999, which is
    the millisecond resolution documented rather than refused."""
    path = tmp_path / "out.xlsx"
    rows = [{"clock": datetime.time(23, 59, 59, us)} for us in (999000, 999499)]
    write_xlsx(str(path), rows, table_name="rows")

    last = datetime.time(23, 59, 59, 999000)
    assert _rows(path) == [{"clock": last}, {"clock": last}]
    # The package reader has no time-only cell type either, so a time reads back as
    # that time on Excel's day zero, 1899-12-31.
    day_zero = datetime.datetime.combine(datetime.date(1899, 12, 31), last)
    assert _read_via_source(path) == [{"clock": day_zero}, {"clock": day_zero}]


@pytest.mark.parametrize(
    "at",
    [
        datetime.datetime.max.replace(
            tzinfo=datetime.timezone(datetime.timedelta(hours=-12))
        ),
        datetime.datetime.min.replace(
            tzinfo=datetime.timezone(datetime.timedelta(hours=12))
        ),
    ],
    ids=["max-west", "min-east"],
)
def test_an_aware_datetime_outside_the_utc_range_is_refused(tmp_path, at):
    """Converting to UTC overflows, and the error names the column rather than
    surfacing as a bare `OverflowError`."""
    _assert_refused(tmp_path, [{"at": at}], r"Column 'at'.*row 1.*UTC")


def test_an_aware_time_of_day_is_refused(tmp_path):
    clock = datetime.time(9, 30, tzinfo=datetime.timezone.utc)
    _assert_refused(tmp_path, [{"clock": clock}], r"Column 'clock'.*timezone")


def test_more_rows_than_a_worksheet_holds_is_refused(tmp_path, monkeypatch):
    """Past the last row xlsxwriter drops the write and carries on, which would leave
    a truncated file. The limit is lowered here rather than a million rows built."""
    monkeypatch.setattr(writer, "_XLSX_MAX_ROWS", 3)
    path = tmp_path / "out.xlsx"
    write_xlsx(str(path), [{"id": 1}, {"id": 2}], table_name="rows")
    path.unlink()

    _assert_refused(tmp_path, [{"id": i} for i in range(3)], "3 rows")


def test_more_columns_than_a_worksheet_holds_is_refused(tmp_path, monkeypatch):
    monkeypatch.setattr(writer, "_XLSX_MAX_COLUMNS", 2)
    path = tmp_path / "out.xlsx"
    write_xlsx(str(path), [{"a": 1, "b": 2}], table_name="rows")
    path.unlink()

    _assert_refused(tmp_path, [{"a": 1, "b": 2, "c": 3}], "3 columns")


def test_a_refusal_late_in_the_load_leaves_no_file_behind(tmp_path):
    """The check runs as each cell is written, so this is the case that proves the
    path is only created when the workbook is closed."""
    rows = [{"id": i} for i in range(1000)] + [{"id": 2**60}]
    _assert_refused(tmp_path, rows, r"row 1001")


def test_the_writer_without_xlsxwriter_names_the_install(tmp_path, monkeypatch):
    monkeypatch.setitem(sys.modules, "xlsxwriter", None)

    with pytest.raises(
        MissingDecoderError, match=r"pip install 'dlt-filesystem\[spreadsheet\]'"
    ):
        write_xlsx(str(tmp_path / "out.xlsx"), [{"id": 1}], table_name="rows")
